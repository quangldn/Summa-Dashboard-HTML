#!/usr/bin/env python3
"""Pull saved CONFIG sheets out of a GX BOM Configurator as regression fixtures.

Every workbook ships with whatever configuration was last saved in it, and each
CONFIG sheet carries both the inputs and the BOM the macro produced from them.
That is a free, authoritative test case: if the Summa engine reads the same
inputs and produces a different BOM, the Summa engine is wrong.

    python3 tools/extract-fixtures.py "GX BOM Configurator V3.0.xlsm" \
        -o tools/fixtures

Writes one JSON per non-empty CONFIG sheet: {config, expected: [{pon, qty}]}.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

# Sheet cells that hold the chassis-level answers (CONFIG #n layout, V3.0).
CELLS = {
    "chassisType": "C4",
    "name": "C5",
    "variantDesc": "C9",
    "powerCable": "C10",
    "ethernetCable": "C11",
    "mountingKit": "C12",
    "powerOption": "C13",
    "controllerOption": "C14",
    "fillers": "C15",
    "fibers": "C16",
    "dustFilter": "C17",
    "airBaffle": "C18",
    "cableGuide": "C19",
    "release": "C22",
    "swLevel": "C23",
    "encryption": "C24",
    "nmsType": "C27",
    "nmsEnable": "C28",
}

# Values the template writes into a cell that has not been answered.
UNSET = re.compile(
    r"^(UNUSED|EMPTY|NA|N/A|Select .*|Choose .*|Not applicable|"
    r"CLNT\d*|LINE\d*|PORT\d*|TRIB\d*)$", re.I)


def unset(v) -> bool:
    return v is None or v == "" or bool(UNSET.match(str(v).strip()))


def yes(v) -> bool:
    return str(v).strip().lower() == "yes"


def shift(coord: str, drow: int, dcol: int) -> tuple[int, int]:
    col, row = coordinate_from_string(coord)
    return row + drow, column_index_from_string(col) + dcol


def extract(path: Path, out_dir: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    aux = wb["AUX"]

    # geometry: chassis -> slot anchors, slot size, line module size
    def table(name):
        t = aux.tables[name]
        rows = [[c.value for c in r] for r in aux[t.ref]]
        hdr = rows[0]
        return [dict(zip(hdr, r)) for r in rows[1:]]

    sled_pos = {r["CHASSIS TYPE / SLED POSITION"]: r for r in table("TEMPLATE_SLED_POSITION")}
    settings = {r["GENERAL SETTINGS"]: r for r in table("TEMPLATE_SETTINGS")}
    sled_def = {r["SLED"]: r for r in table("SLED_DEFINITION") if r.get("SLED")}
    variants = {}
    for t in aux.tables:
        if t.endswith("_CHASSIS_OPTIONS"):
            for r in table(t):
                key = list(r.values())
                variants[r.get("DESCRIPTION")] = key[0]

    written = []
    for ws in wb.worksheets:
        if not ws.title.startswith("CONFIG #"):
            continue
        vals = {k: ws[c].value for k, c in CELLS.items()}
        ctype = vals["chassisType"]
        if unset(ctype) or ctype == "NONE" or unset(vals["variantDesc"]):
            continue

        geo = sled_pos.get(ctype)
        gen = settings.get(ctype)
        if not geo or not gen:
            continue
        n_slots = int(geo["NSLOTS"])
        slot_size = int(geo["SLOT SIZE"])
        port_size = int(gen["LINE MODULE SIZE"])

        slots = []
        for i in range(1, n_slots + 1):
            anchor = geo.get(f"SLOT#{i}")
            if not anchor or anchor == "NA":
                slots.append({"sled": "EMPTY", "line": {}, "client": {}})
                continue
            srow, scol = shift(anchor, 1, 0)
            sled = ws.cell(srow, scol).value
            rec = {"sled": "EMPTY" if unset(sled) else sled, "line": {}, "client": {}}
            if rec["sled"] != "EMPTY":
                sd = sled_def.get(rec["sled"], {})
                lrow, lcol = shift(anchor, 4, 0)
                for b in range(2):
                    n = int(sd.get(f"LINE PLUGGABLES SLOT {b + 1}") or 0)
                    for j in range(n):
                        v = ws.cell(lrow, lcol + slot_size * b + port_size * j).value
                        if not unset(v):
                            rec["line"][f"L{b}.{j}"] = v
                crow, ccol = shift(anchor, 5, 0)
                for b in range(3):
                    n = int(sd.get(f"CLIENT PLUGGABLES SLOT {b + 1}") or 0)
                    for j in range(n):
                        # cages fill row-major across the slot's own columns
                        v = ws.cell(crow + j // slot_size,
                                    ccol + slot_size * b + j % slot_size).value
                        if not unset(v):
                            rec["client"][f"C{b}.{j}"] = v
            slots.append(rec)

        expected, r = [], 32
        while ws.cell(r, 2).value:
            expected.append({"pon": str(ws.cell(r, 2).value),
                             "description": ws.cell(r, 3).value,
                             "qty": ws.cell(r, 5).value})
            r += 1
        if not expected:
            continue

        cfg = {
            "name": vals["name"] or ws.title,
            "chassisType": ctype,
            "variant": variants.get(vals["variantDesc"]),
            "variantDesc": vals["variantDesc"],
            "powerOption": None if unset(vals["powerOption"]) else vals["powerOption"],
            "controllerOption": None if unset(vals["controllerOption"]) else vals["controllerOption"],
            "mountingKit": None if unset(vals["mountingKit"]) else vals["mountingKit"],
            "powerCable": None if unset(vals["powerCable"]) or
                          str(vals["powerCable"]).lower().startswith("procured")
                          else vals["powerCable"],
            "ethernetCable": None if unset(vals["ethernetCable"]) or
                             str(vals["ethernetCable"]).lower().startswith("procured")
                             else vals["ethernetCable"],
            "dustFilter": yes(vals["dustFilter"]),
            "airBaffle": yes(vals["airBaffle"]),
            "cableGuide": yes(vals["cableGuide"]),
            "fillers": yes(vals["fillers"]),
            "release": None if unset(vals["release"]) else vals["release"],
            "swLevel": None if unset(vals["swLevel"]) else vals["swLevel"],
            "encryption": yes(vals["encryption"]),
            "nmsEnable": yes(vals["nmsEnable"]),
            "nmsType": None if unset(vals["nmsType"]) else vals["nmsType"],
            "l0Restoration": False,
            "l0l1Automation": False,
            "slots": slots,
        }

        out_dir.mkdir(parents=True, exist_ok=True)
        stem = re.sub(r"[^A-Za-z0-9]+", "-", f"{path.stem}-{ws.title}").strip("-").lower()
        dest = out_dir / f"{stem}.json"
        dest.write_text(json.dumps(
            {"source": path.name, "sheet": ws.title, "config": cfg, "expected": expected},
            indent=2), encoding="utf-8")
        written.append((dest, len(expected)))

    for d, n in written:
        print(f"  {d}  ({n} BOM lines)")
    if not written:
        print("  no saved configurations found in this workbook")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=Path("tools/fixtures"))
    a = ap.parse_args()
    extract(a.workbook, a.out)


if __name__ == "__main__":
    main()
